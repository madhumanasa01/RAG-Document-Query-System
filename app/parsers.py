"""Turn uploaded files into page-level text, falling back to OCR for scanned content."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from app.config import settings

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}


@dataclass
class Page:
    number: int
    text: str
    ocr: bool = False


class UnsupportedFileType(ValueError):
    pass


def parse(filename: str, data: bytes) -> list[Page]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise UnsupportedFileType(f"Unsupported file type: {ext or 'none'}")
    if ext == ".pdf":
        return _parse_pdf(data)
    if ext == ".docx":
        return _parse_docx(data)
    if ext in {".txt", ".md"}:
        return [Page(1, data.decode("utf-8", errors="ignore"))]
    return [Page(1, _ocr_image_bytes(data), ocr=True)]


def _parse_pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if len(text) < settings.ocr_min_chars_per_page:
            # Little or no text layer: treat it as a scanned page.
            ocr_text = _ocr_pdf_page(data, i)
            if ocr_text:
                pages.append(Page(i, ocr_text, ocr=True))
                continue
        pages.append(Page(i, text))
    return pages


def _parse_docx(data: bytes) -> list[Page]:
    import docx

    document = docx.Document(BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Page(1, "\n".join(parts))]


def _ocr_image_bytes(data: bytes) -> str:
    import pytesseract
    from PIL import Image

    return pytesseract.image_to_string(Image.open(BytesIO(data)), lang=settings.ocr_lang).strip()


def _ocr_pdf_page(data: bytes, page_number: int) -> str:
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError:
        return ""
    pdf = pdfium.PdfDocument(data)
    image = pdf[page_number - 1].render(scale=300 / 72).to_pil()
    return pytesseract.image_to_string(image, lang=settings.ocr_lang).strip()
